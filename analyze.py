#!/usr/bin/env python3
"""ChannelTalk CS complaint analyzer — produces a formatted Excel report."""

import argparse
import os
import sys
from datetime import datetime

import anthropic
import pandas as pd
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

load_dotenv()

CONSULTANT_ORDER = ["무민", "쵸비", "준타", "솔트", "라미", "테디", "펑키", "알프"]

COLORS = [
    "FFD700", "90EE90", "ADD8E6", "FFB6C1",
    "FFA07A", "E6E6FA", "98FB98", "87CEEB", "DDA0DD", "F0E68C"
]


def find_sheet(sheets: dict, *keywords) -> pd.DataFrame | None:
    for name, df in sheets.items():
        name_lower = name.lower()
        if any(kw.lower() in name_lower for kw in keywords):
            return df
    return None


def load_excel(filepath: str) -> dict[str, pd.DataFrame]:
    return pd.read_excel(filepath, sheet_name=None, dtype=str)


def format_phone(raw) -> str:
    if pd.isna(raw) or str(raw).strip() in ("", "nan", "None"):
        return ""
    phone = str(raw).strip().replace("-", "").replace(" ", "")
    if phone.startswith("8210"):
        phone = "0" + phone[2:]
    elif phone.startswith("82"):
        phone = "0" + phone[2:]
    return phone


def format_date_range(created_at, closed_at) -> str:
    fmt = "%m-%d"
    try:
        start = pd.to_datetime(created_at)
        end = pd.to_datetime(closed_at)
        if pd.isna(start):
            return ""
        s = start.strftime(fmt)
        if pd.isna(end) or start.date() == end.date():
            return s
        return f"{s}~{end.strftime(fmt)}"
    except Exception:
        return str(created_at) if not pd.isna(created_at) else ""


def calc_tat(created_at, closed_at) -> int | str:
    try:
        start = pd.to_datetime(created_at)
        end = pd.to_datetime(closed_at)
        if pd.isna(start) or pd.isna(end):
            return ""
        delta = end - start
        return max(0, int(delta.total_seconds() // 60))
    except Exception:
        return ""


def is_private(val) -> bool:
    if isinstance(val, bool):
        return val
    if isinstance(val, (int, float)):
        return bool(val)
    return str(val).strip().lower() in ("true", "1", "yes")


def get_latest_time(messages_df: pd.DataFrame, chat_id: str):
    chat_msgs = messages_df[messages_df["chatId"] == chat_id].copy()
    if chat_msgs.empty:
        return pd.NaT

    times = []

    # Last internal (isPrivate=True) message time
    internal = chat_msgs[chat_msgs["isPrivate"].apply(is_private)]
    if not internal.empty:
        t = pd.to_datetime(internal["createdAt"], errors="coerce").max()
        if not pd.isna(t):
            times.append(t)

    # Last manager public message time
    manager_public = chat_msgs[
        (~chat_msgs["isPrivate"].apply(is_private)) &
        (chat_msgs["personType"].str.lower().isin(["manager", "bot"]))
    ]
    if not manager_public.empty:
        t = pd.to_datetime(manager_public["createdAt"], errors="coerce").max()
        if not pd.isna(t):
            times.append(t)

    # Last user utterance time
    user_msgs = chat_msgs[chat_msgs["personType"].str.lower() == "user"]
    if not user_msgs.empty:
        t = pd.to_datetime(user_msgs["createdAt"], errors="coerce").max()
        if not pd.isna(t):
            times.append(t)

    return max(times) if times else pd.NaT


def get_consultants(messages_df: pd.DataFrame, managers_df: pd.DataFrame, chat_id: str) -> str:
    chat_msgs = messages_df[messages_df["chatId"] == chat_id]
    public_manager = chat_msgs[
        (~chat_msgs["isPrivate"].apply(is_private)) &
        (chat_msgs["personType"].str.lower() == "manager")
    ]
    if public_manager.empty:
        return "알프"

    seen = []
    for person_id in public_manager["personId"].dropna():
        person_id = str(person_id).strip()
        if not managers_df.empty:
            match = managers_df[managers_df["id"].astype(str).str.strip() == person_id]
            if not match.empty:
                name = str(match.iloc[0].get("name", "")).strip()
                if name and name not in seen:
                    seen.append(name)
                continue
        if person_id not in seen:
            seen.append(person_id)

    return ", ".join(seen) if seen else "알프"


def get_user_content(messages_df: pd.DataFrame, chat_id: str) -> str:
    chat_msgs = messages_df[messages_df["chatId"] == chat_id].copy()
    user_msgs = chat_msgs[
        (chat_msgs["personType"].str.lower() == "user") &
        (~chat_msgs["isPrivate"].apply(is_private))
    ].copy()

    if user_msgs.empty:
        return ""

    user_msgs["createdAt"] = pd.to_datetime(user_msgs["createdAt"], errors="coerce")
    user_msgs = user_msgs.sort_values("createdAt")

    parts = []
    prev_sender = None
    buffer = []

    for _, row in user_msgs.iterrows():
        sender = str(row.get("personId", "")).strip()
        text = str(row.get("plainText", row.get("text", ""))).strip()
        if not text or text in ("nan", "None"):
            continue
        if sender == prev_sender:
            buffer.append(text)
        else:
            if buffer:
                parts.append(" ".join(buffer))
            buffer = [text]
            prev_sender = sender

    if buffer:
        parts.append(" ".join(buffer))

    return "\n".join(parts)


def summarize_content(client: anthropic.Anthropic, content: str) -> str:
    if not content.strip():
        return ""
    response = client.messages.create(
        model="claude-haiku-4-5",
        max_tokens=200,
        messages=[{
            "role": "user",
            "content": (
                "다음 고객 상담 내용을 1~2문장으로 핵심만 간단히 요약해주세요. "
                "한국어로 작성하고, 불필요한 설명 없이 요약만 출력하세요.\n\n"
                f"상담 내용:\n{content}"
            )
        }]
    )
    return response.content[0].text.strip()


def get_primary_consultant(messages_df: pd.DataFrame, managers_df: pd.DataFrame, chat_id: str) -> str:
    chat_msgs = messages_df[messages_df["chatId"] == chat_id]
    public_manager = chat_msgs[
        (~chat_msgs["isPrivate"].apply(is_private)) &
        (chat_msgs["personType"].str.lower() == "manager")
    ]
    if public_manager.empty:
        return "알프"

    counts = public_manager["personId"].value_counts()
    if counts.empty:
        return "알프"

    top_id = str(counts.index[0]).strip()
    if not managers_df.empty:
        match = managers_df[managers_df["id"].astype(str).str.strip() == top_id]
        if not match.empty:
            return str(match.iloc[0].get("name", top_id)).strip()
    return top_id


def get_nested(row, *keys):
    for key in keys:
        val = row.get(key)
        if val is not None and str(val).strip() not in ("", "nan", "None"):
            return str(val).strip()
    return ""


def analyze(filepath: str, output_path: str, on_progress=None):
    """Run the full analysis pipeline.

    on_progress(current, total, message) is called after each consultation is processed.
    """
    api_key = os.getenv("ANTHROPIC_API_KEY")
    if not api_key:
        raise ValueError("ANTHROPIC_API_KEY not set. Copy .env.example to .env and add your key.")

    client = anthropic.Anthropic(api_key=api_key)

    print(f"Loading {filepath} ...")
    sheets = load_excel(filepath)

    userchat_df = find_sheet(sheets, "userchat", "UserChat", "chat")
    message_df = find_sheet(sheets, "message", "Message")
    manager_df = find_sheet(sheets, "manager", "Manager")
    user_df = find_sheet(sheets, "user", "User")

    if userchat_df is None:
        raise ValueError("Could not find UserChat sheet in the Excel file.")
    if message_df is None:
        raise ValueError("Could not find Message sheet in the Excel file.")

    # Normalize column names
    userchat_df.columns = [c.strip() for c in userchat_df.columns]
    message_df.columns = [c.strip() for c in message_df.columns]
    if manager_df is not None:
        manager_df.columns = [c.strip() for c in manager_df.columns]
    else:
        manager_df = pd.DataFrame(columns=["id", "name"])
    if user_df is not None:
        user_df.columns = [c.strip() for c in user_df.columns]
    else:
        user_df = pd.DataFrame()

    # Ensure required columns exist in message_df
    for col in ["chatId", "personType", "personId", "isPrivate", "createdAt", "plainText"]:
        if col not in message_df.columns:
            message_df[col] = ""

    # Build rows
    rows = []
    total = len(userchat_df)
    print(f"Processing {total} consultations...")

    for idx, chat in userchat_df.iterrows():
        chat_id = str(chat.get("id", "")).strip()
        seq = idx + 1

        # Sort key
        sort_time = get_latest_time(message_df, chat_id)

        # Time range
        created_at = chat.get("createdAt", "")
        closed_at = chat.get("closedAt", "")
        time_str = format_date_range(created_at, closed_at)

        # Consultants
        consultants = get_consultants(message_df, manager_df, chat_id)

        # User name
        user_id = str(chat.get("userId", "")).strip()
        user_name = ""
        if not user_df.empty and user_id:
            u_match = user_df[user_df["id"].astype(str).str.strip() == user_id]
            if not u_match.empty:
                user_name = get_nested(u_match.iloc[0], "name", "profile.name")
        if not user_name:
            user_name = get_nested(chat, "userName", "name")

        # Phone
        phone_raw = ""
        if not user_df.empty and user_id:
            u_match = user_df[user_df["id"].astype(str).str.strip() == user_id]
            if not u_match.empty:
                phone_raw = get_nested(u_match.iloc[0], "mobileNumber", "profile.mobileNumber", "phoneNumber")
        phone = format_phone(phone_raw)

        # Re-claim and 이탈여부 — from chat tags or custom fields; leave blank if not present
        reclaim = get_nested(chat, "reclaim", "reClaim", "재클레임")
        churn = get_nested(chat, "churn", "이탈여부")

        # Content
        content = get_user_content(message_df, chat_id)

        # Summary
        print(f"  [{seq}/{total}] Summarizing chat {chat_id}...")
        summary = summarize_content(client, content)

        if on_progress:
            on_progress(seq, total, f"[{seq}/{total}] {chat_id} 요약 완료")

        # CSAT
        csat = get_nested(chat, "csat", "profile.csat", "CSAT")
        csat_reason = get_nested(chat, "csatComment", "profile.csatComment", "이유")

        # TAT
        tat = calc_tat(created_at, closed_at)

        # Primary consultant for Sheet2
        primary = get_primary_consultant(message_df, manager_df, chat_id)

        rows.append({
            "sort_time": sort_time,
            "상담 번호": seq,
            "시간": time_str,
            "상담사": consultants,
            "유저": user_name,
            "연락처": phone,
            "재클레임": reclaim,
            "이탈여부": churn,
            "요약": summary,
            "상담 내용": content,
            "CSAT": csat,
            "이유": csat_reason,
            "TAT": tat,
            "상담 ID": chat_id,
            "_primary": primary,
        })

    # Sort by sort_time ascending
    rows.sort(key=lambda r: (r["sort_time"] is pd.NaT, r["sort_time"]))

    # Re-number after sort
    for i, row in enumerate(rows, 1):
        row["상담 번호"] = i

    # ── Write Excel ──────────────────────────────────────────────────────────
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "상담 내역"

    sheet1_cols = [
        "상담 번호", "시간", "상담사", "유저", "연락처",
        "재클레임", "이탈여부", "요약", "상담 내용",
        "CSAT", "이유", "TAT", "상담 ID"
    ]

    base_font = Font(name="Arial", size=10)
    header_font = Font(name="Arial", size=10, bold=True)
    no_wrap = Alignment(wrap_text=False)

    # Header
    ws1.append(sheet1_cols)
    for cell in ws1[1]:
        cell.font = header_font
        cell.alignment = no_wrap

    # Freeze header
    ws1.freeze_panes = "A2"

    # Detect duplicate user names for color coding
    name_counts: dict[str, int] = {}
    for row in rows:
        n = row["유저"]
        name_counts[n] = name_counts.get(n, 0) + 1
    dup_names = {n for n, c in name_counts.items() if c > 1 and n}

    name_color_map: dict[str, str] = {}
    color_idx = 0
    for row in rows:
        n = row["유저"]
        if n in dup_names and n not in name_color_map:
            name_color_map[n] = COLORS[color_idx % len(COLORS)]
            color_idx += 1

    # Data rows
    for row in rows:
        ws1.append([row[c] for c in sheet1_cols])
        excel_row = ws1.max_row
        name = row["유저"]
        fill = PatternFill("solid", fgColor=name_color_map[name]) if name in name_color_map else None
        for col_idx, cell in enumerate(ws1[excel_row], 1):
            cell.font = base_font
            cell.alignment = no_wrap
            if fill:
                cell.fill = fill

    # Column widths (approximate)
    col_widths = {
        "상담 번호": 8, "시간": 14, "상담사": 14, "유저": 12, "연락처": 14,
        "재클레임": 10, "이탈여부": 10, "요약": 40, "상담 내용": 60,
        "CSAT": 8, "이유": 20, "TAT": 8, "상담 ID": 24,
    }
    for i, col_name in enumerate(sheet1_cols, 1):
        ws1.column_dimensions[get_column_letter(i)].width = col_widths.get(col_name, 12)

    # ── Sheet2: 상담사별 집계 ─────────────────────────────────────────────────
    ws2 = wb.create_sheet("상담사별 집계")

    consultant_counts: dict[str, int] = {}
    for row in rows:
        c = row["_primary"]
        consultant_counts[c] = consultant_counts.get(c, 0) + 1

    ordered = [c for c in CONSULTANT_ORDER if c in consultant_counts]
    others = [c for c in consultant_counts if c not in CONSULTANT_ORDER]
    ordered += sorted(others)

    ws2.append(["상담사", "상담 건수"])
    for cell in ws2[1]:
        cell.font = header_font
        cell.alignment = no_wrap

    total_count = 0
    for name in ordered:
        count = consultant_counts.get(name, 0)
        total_count += count
        row_data = [name, count]
        ws2.append(row_data)
        for cell in ws2[ws2.max_row]:
            cell.font = base_font
            cell.alignment = no_wrap

    ws2.append(["합계", total_count])
    for cell in ws2[ws2.max_row]:
        cell.font = Font(name="Arial", size=10, bold=True)
        cell.alignment = no_wrap

    ws2.column_dimensions["A"].width = 14
    ws2.column_dimensions["B"].width = 12

    wb.save(output_path)
    print(f"\nDone! Output saved to: {output_path}")


def main():
    parser = argparse.ArgumentParser(description="ChannelTalk CS complaint analyzer")
    parser.add_argument("input", help="Path to ChannelTalk Excel export (.xlsx)")
    parser.add_argument("-o", "--output", help="Output file path (default: <input>_analyzed.xlsx)")
    args = parser.parse_args()

    if not os.path.exists(args.input):
        print(f"Error: File not found: {args.input}", file=sys.stderr)
        sys.exit(1)

    if args.output:
        output_path = args.output
    else:
        base = os.path.splitext(args.input)[0]
        output_path = f"{base}_analyzed.xlsx"

    analyze(args.input, output_path)


if __name__ == "__main__":
    main()
